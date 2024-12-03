import json
from docx import Document
import os
import shutil

# スクリプトのディレクトリ
script_dir = os.path.dirname(os.path.abspath(__file__))

# JSONファイルの絶対パスを指定
json_path = os.path.join(script_dir, 'data.json')

# テンプレートWord文書のパス
template_path = os.path.join(script_dir, 'temp.docx')  # ここを変更
# 出力先Word文書のパス
output_word_path = os.path.join(script_dir, 'houkokusyo.docx')  # ここを変更

# テンプレートを出力先にコピーして初期化
try:
    shutil.copyfile(template_path, output_word_path)
    print("テンプレートのコピー成功")
except FileNotFoundError as e:
    print(f"Error: テンプレートファイルが見つかりません: {e}")
    exit()

try:
    # JSONファイルを読み込み
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    print("JSONファイルの読み込み成功")
    print(f"JSONファイルの内容: {data}")
except FileNotFoundError:
    print("Error: data.json が見つかりません。")
    exit()
except json.JSONDecodeError:
    print("Error: data.json の解析に失敗しました。")
    exit()

# テンプレートWord文書を開く
print(f"Word文書パス: {output_word_path}")

try:
    doc = Document(output_word_path)
    print("Word文書の読み込み成功")
except Exception as e:
    print(f"Word文書の読み込みに失敗しました: {e}")
    exit()

# テーブルを探し、データを特定の欄に追加
try:
    table = doc.tables[0]  # 文書内の最初のテーブルを取得

    # テーブルのセルにデータを追加
    if '経緯と状況' in data:
        table.cell(3, 1).text = '\n'.join(data['経緯と状況'])
    if '要因' in data:
        table.cell(4, 1).text = '\n'.join(data['要因'])
    if '対策' in data:
        table.cell(5, 1).text = '\n'.join(data['対策'])

    print("データの追加成功")
except IndexError as e:
    print(f"Error: テーブルのインデックスが見つかりません: {e}")
except KeyError as e:
    print(f"Error: キーが見つかりません: {e}")

# Word文書を保存
try:
    doc.save(output_word_path)
    print(f"Word文書が {output_word_path} に保存されました。")
except Exception as e:
    print(f"Word文書の保存に失敗しました: {e}")
